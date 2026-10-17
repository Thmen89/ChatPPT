"""The in-process PPTX writer (services/pptx_writer.py) and its endpoint."""

from __future__ import annotations

import asyncio
import base64
import io
import zipfile
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

from api.v1.ppt.endpoints.export import EXPORT_ROUTER
from models.pptx_models import PptxObjectFitModel, PptxPresentationModel
from services import pptx_writer


def _png_data_uri(width: int = 40, height: int = 20, color: tuple[int, int, int, int] = (200, 30, 30, 255)) -> str:
    buffer = io.BytesIO()
    Image.new("RGBA", (width, height), color).save(buffer, format="PNG")
    return "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode()


def _deck(name: str = "Quarterly review: Q3") -> dict:
    font = {"name": "Inter", "size": 24, "font_weight": 700, "italic": False, "color": "112233"}
    return {
        "name": name,
        "slides": [
            {
                "background": {"color": "f0f0f0", "opacity": 1},
                "note": "Say hello",
                "shapes": [
                    {
                        "shape_type": "textbox",
                        "position": {"left": 100, "top": 40, "width": 600, "height": 80},
                        "margin": {"top": 4, "bottom": 4, "left": 8, "right": 8},
                        "vertical_alignment": 3,
                        "paragraphs": [
                            {
                                "alignment": 2,
                                "font": font,
                                "line_height": 30,
                                "text": "Hello world",
                                "text_runs": [
                                    {"text": "Hello ", "font": font},
                                    {"text": "world", "font": {**font, "italic": True, "font_weight": 400}},
                                ],
                            }
                        ],
                    },
                    {
                        "shape_type": "textbox",
                        "position": {"left": 84, "top": 200, "width": 500, "height": 120},
                        "paragraphs": [
                            {
                                "text": "First",
                                "font": font,
                                "level": 0,
                                "bullet": {"type": "default", "color": "ff0000"},
                            },
                            {"text": "Second", "font": font, "level": 0, "bullet": {"type": "default"}},
                        ],
                    },
                    {
                        "shape_type": "autoshape",
                        "type": 5,
                        "position": {"left": 700, "top": 200, "width": 300, "height": 200},
                        "fill": {"color": "336699", "opacity": 0.5},
                        "stroke": {"color": "000000", "thickness": 2, "opacity": 1},
                        "shadow": {"radius": 8, "offset": 4, "color": "000000", "opacity": 0.3, "angle": 90},
                        "border_radius": 16,
                    },
                    {
                        "shape_type": "connector",
                        "type": 1,
                        "position": {"left": 100, "top": 500, "width": 1000, "height": 0},
                        "thickness": 2,
                        "color": "999999",
                    },
                    {
                        "shape_type": "picture",
                        "position": {"left": 1000, "top": 500, "width": 40, "height": 20},
                        "clip": False,
                        "shape": "rectangle",
                        "object_fit": {"fit": "fill"},
                        "picture": {"is_network": False, "path": _png_data_uri()},
                    },
                    {
                        "shape_type": "picture",
                        "position": {"left": 10, "top": 10, "width": 40, "height": 40},
                        "object_fit": {"fit": "contain"},
                        "picture": {"is_network": False, "path": "/nowhere/missing.png"},
                    },
                ],
            },
            {"shapes": []},
        ],
    }


@pytest.fixture()
def data(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setenv("APP_DATA_DIRECTORY", str(tmp_path / "data"))
    monkeypatch.setenv("TEMP_DIRECTORY", str(tmp_path / "tmp"))
    monkeypatch.setenv("DISABLE_AUTH", "true")
    (tmp_path / "data").mkdir()
    return tmp_path / "data"


def _write(model: dict) -> Path:
    return Path(asyncio.run(pptx_writer.write_pptx_to_exports(PptxPresentationModel.model_validate(model))))


def test_writes_a_deck_python_pptx_reads_back(data: Path) -> None:
    path = _write(_deck())

    assert path.parent == data / "exports"
    assert path.name.startswith("Quarterly-review-Q3_") and path.suffix == ".pptx"
    assert path.stat().st_mode & 0o777 == 0o644
    prs = Presentation(str(path))
    assert len(prs.slides) == 2
    assert (Emu(prs.slide_width or 0).pt, Emu(prs.slide_height or 0).pt) == (960, 540)

    slide = prs.slides[0]
    kinds = [shape.shape_type for shape in slide.shapes]
    assert [int(k) for k in kinds] == [17, 17, 1, 9, 13]  # text, text, autoshape, line, picture
    title = slide.shapes[0]
    assert (Emu(title.left).pt, Emu(title.top).pt, Emu(title.width).pt) == (75, 30, 450)
    assert title.text_frame.text == "Hello world"
    runs = title.text_frame.paragraphs[0].runs
    assert [r.text for r in runs] == ["Hello ", "world"]
    assert runs[0].font.bold is True and runs[1].font.italic is True
    assert str(runs[0].font.color.rgb) == "112233"
    assert [p.text for p in slide.shapes[1].text_frame.paragraphs] == ["  First", "  Second"]
    assert (
        slide.shapes[1]
        .text_frame.paragraphs[0]
        ._p.pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar")
        is not None
    )
    assert str(slide.shapes[2].fill.fore_color.rgb) == "336699"
    assert slide.shapes[4].image.size == (40, 20)
    assert slide.has_notes_slide and slide.notes_slide.notes_text_frame.text == "Say hello"
    assert str(slide.background.fill.fore_color.rgb) == "F0F0F0"

    with zipfile.ZipFile(path) as package:
        assert package.testzip() is None
        assert "docProps/thumbnail.jpeg" not in package.namelist()
        assert b"notesMasterIdLst" in package.read("ppt/presentation.xml")


def test_unknown_object_fit_stretches_instead_of_failing() -> None:
    assert PptxObjectFitModel.model_validate({"fit": "scale-down"}).fit is None
    assert PptxObjectFitModel.model_validate({"fit": "cover"}).fit is not None


def test_export_basename_follows_the_exporter_rule() -> None:
    assert pptx_writer.export_basename("  My deck: v2 ") == "My-deck-v2"
    assert pptx_writer.export_basename("ÄÖÜ") == "presentation"
    assert pptx_writer.export_basename(None) == "presentation"


def test_picture_sources_resolve_to_local_files(data: Path) -> None:
    (data / "images").mkdir()
    Image.new("RGB", (8, 8)).save(data / "images" / "a.png")
    model = PptxPresentationModel.model_validate(
        {
            "slides": [
                {
                    "shapes": [
                        {"shape_type": "picture", "position": {}, "picture": {"path": "/app_data/images/a.png"}},
                        {"shape_type": "picture", "position": {}, "picture": {"path": _png_data_uri()}},
                        {"shape_type": "picture", "position": {}, "picture": {"path": "/etc/passwd"}},
                        {"shape_type": "picture", "position": {}, "picture": {"path": "data:image/png;base64,"}},
                    ]
                }
            ]
        }
    )
    temp = data.parent / "resolve"
    temp.mkdir()
    asyncio.run(pptx_writer.resolve_picture_sources(model, str(temp)))
    paths = [shape.picture.path for shape in model.slides[0].shapes]  # type: ignore[union-attr]
    assert paths[0] == str(data / "images" / "a.png")
    assert paths[1].startswith(str(temp)) and Path(paths[1]).is_file()
    assert paths[2:] == ["", ""]


def test_endpoint_writes_into_exports(data: Path) -> None:
    app = FastAPI()
    app.include_router(EXPORT_ROUTER, prefix="/api/v1/ppt")
    response = TestClient(app).post("/api/v1/ppt/export/pptx", json=_deck("Deck"))
    assert response.status_code == 200, response.text
    path = Path(response.json()["path"])
    assert path.is_file() and path.parent == data / "exports"
    assert (
        TestClient(app)
        .post("/api/v1/ppt/export/pptx", json={"slides": [{"shapes": [{"shape_type": "x"}]}]})
        .status_code
        == 422
    )
