"""Write a .pptx from the exporter's shape models (models/pptx_models.py), in-process.

The exporter (Next.js, headless Chromium) measures every slide and sends positioned
text boxes, shapes, pictures and lines; this turns them into python-pptx objects. It
follows the behaviour of the converter that shipped as a PyInstaller binary in the
earlier export bundle, so the files match, without a process start per export.
"""

from __future__ import annotations

import asyncio
import base64
import binascii
import io
import logging
import os
import re
import tempfile
import uuid
import zipfile
from typing import Iterable, List, Optional
from urllib.parse import urlsplit

import httpx
from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.shapes.autoshape import Shape
from pptx.slide import Slide
from pptx.text.text import Font, TextFrame, _Paragraph, _Run
from pptx.util import Pt

from models.pptx_models import (
    PptxAutoShapeBoxModel,
    PptxBoxShapeEnum,
    PptxConnectorModel,
    PptxFillModel,
    PptxFontModel,
    PptxObjectFitEnum,
    PptxObjectFitModel,
    PptxParagraphBulletModel,
    PptxParagraphModel,
    PptxPictureBoxModel,
    PptxPositionModel,
    PptxPresentationModel,
    PptxShadowModel,
    PptxSlideModel,
    PptxSpacingModel,
    PptxStrokeModel,
    PptxTextBoxModel,
    PptxTextRunModel,
    px_to_pt,
)

LOGGER = logging.getLogger(__name__)

BLANK_SLIDE_LAYOUT = 6
SLIDE_WIDTH_PX = 1280
SLIDE_HEIGHT_PX = 720
BULLET_CHAR = "•"
BULLET_PREFIXES = (f"{BULLET_CHAR} ", "- ", "* ")
RTL_RANGES = (
    (0x0590, 0x05FF),
    (0x0600, 0x06FF),
    (0x0700, 0x074F),
    (0x0750, 0x077F),
    (0x0780, 0x07BF),
    (0x08A0, 0x08FF),
    (0xFB1D, 0xFB4F),
    (0xFB50, 0xFDFF),
    (0xFE70, 0xFEFF),
    (0x1EE00, 0x1EEFF),
)
DOWNLOAD_TIMEOUT_S = 20.0

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NOTES_MASTER_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesMaster"
THUMBNAIL_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/metadata/thumbnail"
THUMBNAIL_PART = "docProps/thumbnail.jpeg"
# Children of a:pPr that come after the bullet group, in schema order.
_AFTER_BULLETS = ("tabLst", "defRPr", "extLst")


def gtz(value: float) -> float:
    return max(value, 0)


def export_basename(name: Optional[str]) -> str:
    """The exporter's file-name rule: spaces to dashes, ASCII letters, digits and dashes only."""
    cleaned = re.sub(r"[^a-zA-Z0-9-]", "", re.sub(r"\s+", "-", (name or "").strip()))
    return cleaned or "presentation"


def text_has_rtl(text: Optional[str]) -> bool:
    if not text:
        return False
    return any(start <= ord(ch) <= end for ch in text for start, end in RTL_RANGES)


# --------------------------------------------------------------------------- images


def _empty_image() -> Image.Image:
    return Image.new("RGBA", (1, 1), (0, 0, 0, 0))


def clip_image(image: Image.Image, width: int, height: int, focus_x: float = 50.0, focus_y: float = 50.0) -> Image.Image:
    if width <= 0 or height <= 0:
        return _empty_image()
    img_width, img_height = image.size
    img_aspect = img_width / img_height
    if img_aspect > width / height:
        new_height = height
        new_width = int(new_height * img_aspect)
    else:
        new_width = width
        new_height = int(new_width / img_aspect)
    resized = image.resize((new_width, new_height), Image.Resampling.LANCZOS)
    focus_x = max(0.0, min(100.0, focus_x))
    focus_y = max(0.0, min(100.0, focus_y))
    left = int((new_width - width) * (focus_x / 100.0))
    top = int((new_height - height) * (focus_y / 100.0))
    return resized.crop((left, top, left + width, top + height))


def fit_image(image: Image.Image, width: int, height: int, object_fit: PptxObjectFitModel) -> Image.Image:
    if width <= 0 or height <= 0:
        return _empty_image()
    if not object_fit.fit:
        return image
    img_width, img_height = image.size
    img_aspect = img_width / img_height
    box_aspect = width / height
    focus_x, focus_y = 50.0, 50.0
    if object_fit.focus and len(object_fit.focus) == 2:
        focus_x = object_fit.focus[0] if object_fit.focus[0] is not None else 50.0
        focus_y = object_fit.focus[1] if object_fit.focus[1] is not None else 50.0

    if object_fit.fit == PptxObjectFitEnum.CONTAIN:
        if img_aspect > box_aspect:
            new_width, new_height = width, int(width / img_aspect)
        else:
            new_width, new_height = int(height * img_aspect), height
        resized = image.resize((new_width, new_height), Image.Resampling.LANCZOS)
        result = Image.new("RGBA", (width, height), (0, 0, 0, 0))
        result.paste(
            resized,
            (int((width - new_width) * (focus_x / 100.0)), int((height - new_height) * (focus_y / 100.0))),
        )
        return result
    if object_fit.fit == PptxObjectFitEnum.COVER:
        if img_aspect > box_aspect:
            new_width, new_height = int(height * img_aspect), height
        else:
            new_width, new_height = width, int(width / img_aspect)
        resized = image.resize((new_width, new_height), Image.Resampling.LANCZOS)
        left = int((new_width - width) * (focus_x / 100.0))
        top = int((new_height - height) * (focus_y / 100.0))
        return resized.crop((left, top, left + width, top + height))
    if object_fit.fit == PptxObjectFitEnum.FILL:
        if image.size == (width, height):
            return image
        return image.resize((width, height), Image.Resampling.LANCZOS)
    return image


def round_image_corners(image: Image.Image, radii: List[int]) -> Image.Image:
    if len(radii) != 4:
        raise ValueError("Image Border Radius - radii must contain exactly 4 values for each corner")
    w, h = image.size
    max_radius = min(w // 2, h // 2)
    radii = [min(radius, max_radius) for radius in radii]
    if image.mode != "RGBA":
        image = image.convert("RGBA")
    rounded_mask = Image.new("L", image.size, 0)
    rectangular_mask = Image.new("L", image.size, 255)
    for corner, radius in enumerate(radii):
        if radius <= 0:
            continue
        circle = Image.new("L", (radius * 2, radius * 2), 0)
        ImageDraw.Draw(circle).ellipse((0, 0, radius * 2 - 1, radius * 2 - 1), fill=255)
        if corner == 0:
            rounded_mask.paste(circle.crop((0, 0, radius, radius)), (0, 0))
            rectangular_mask.paste(0, (0, 0, radius, radius))
        elif corner == 1:
            rounded_mask.paste(circle.crop((radius, 0, radius * 2, radius)), (w - radius, 0))
            rectangular_mask.paste(0, (w - radius, 0, w, radius))
        elif corner == 2:
            rounded_mask.paste(circle.crop((radius, radius, radius * 2, radius * 2)), (w - radius, h - radius))
            rectangular_mask.paste(0, (w - radius, h - radius, w, h))
        else:
            rounded_mask.paste(circle.crop((0, radius, radius, radius * 2)), (0, h - radius))
            rectangular_mask.paste(0, (0, h - radius, radius, h))
    corner_mask = Image.composite(rounded_mask, rectangular_mask, rounded_mask)
    final_alpha = Image.composite(image.getchannel("A"), Image.new("L", image.size, 0), corner_mask)
    result = Image.new("RGBA", image.size)
    result.paste(image.convert("RGB"), (0, 0))
    result.putalpha(final_alpha)
    return result


def invert_image(image: Image.Image) -> Image.Image:
    rgba = image.convert("RGBA")
    r, g, b, a = rgba.split()
    inverted = Image.merge("RGBA", (*(c.point(lambda v: 255 - v) for c in (r, g, b)), a))
    # Fully transparent pixels become transparent black, as before.
    return Image.composite(inverted, Image.new("RGBA", rgba.size, (0, 0, 0, 0)), a.point(lambda v: 255 if v else 0))


def create_circle_image(image: Image.Image) -> Image.Image:
    img = image.convert("RGBA")
    size = img.size
    radius = min(size) // 2
    cx, cy = size[0] // 2, size[1] // 2
    mask = Image.new("RGBA", size, color=(0, 0, 0, 0))
    ImageDraw.Draw(mask).ellipse((cx - radius, cy - radius, cx + radius, cy + radius), fill=(255, 255, 255, 255))
    return Image.composite(img, mask, mask)


def set_image_opacity(image: Image.Image, opacity: float) -> Image.Image:
    opacity = max(0.0, min(1.0, opacity))
    if image.mode != "RGBA":
        image = image.convert("RGBA")
    new_alpha = image.getchannel("A").point(lambda x: int(x * opacity))
    result = Image.new("RGBA", image.size)
    result.paste(image.convert("RGB"), (0, 0))
    result.putalpha(new_alpha)
    return result


# --------------------------------------------------------------------------- picture sources


def _decode_data_uri(uri: str, temp_dir: str) -> Optional[str]:
    header, _, payload = uri.partition(",")
    if not payload:
        return None
    try:
        data = base64.b64decode(payload) if header.endswith(";base64") else payload.encode()
    except (binascii.Error, ValueError):
        return None
    path = os.path.join(temp_dir, f"{uuid.uuid4()}.img")
    with open(path, "wb") as f:
        f.write(data)
    return path


def _own_origins() -> set[str]:
    """host:port of the servers in this container: what the exporter's page loaded from."""
    from utils.internal_urls import nextjs_base_url

    origins = set()
    for url in (nextjs_base_url(), os.getenv("FAST_API_INTERNAL_URL") or "http://127.0.0.1:8000"):
        parts = urlsplit(url)
        if parts.hostname:
            origins.add(f"{parts.hostname.lower()}:{parts.port or (443 if parts.scheme == 'https' else 80)}")
    return origins


async def _download(client: httpx.AsyncClient, url: str, temp_dir: str, own: set[str]) -> Optional[str]:
    parts = urlsplit(url)
    host_port = f"{(parts.hostname or '').lower()}:{parts.port or (443 if parts.scheme == 'https' else 80)}"
    try:
        if host_port not in own:
            from chatppt.guards import check_outbound_url

            await check_outbound_url(url)
        response = await client.get(url)
        response.raise_for_status()
    except Exception as exc:  # an image that cannot be fetched is left out, as before
        LOGGER.warning("pptx: could not fetch %s: %s", url, exc)
        return None
    path = os.path.join(temp_dir, f"{uuid.uuid4()}.img")
    with open(path, "wb") as f:
        f.write(response.content)
    return path


async def resolve_picture_sources(model: PptxPresentationModel, temp_dir: str) -> None:
    """Point every picture at a local file: data: URIs are decoded, files this app serves
    (/app_data, /static) are read from disk, other URLs are fetched. Pictures whose source
    cannot be read get an empty path and are skipped."""
    from utils.asset_directory_utils import resolve_app_path_to_filesystem

    pictures = [s for slide in model.slides for s in slide.shapes if isinstance(s, PptxPictureBoxModel)]
    to_fetch: dict[str, list[PptxPictureBoxModel]] = {}
    for picture in pictures:
        source = picture.picture.path or ""
        if source.startswith("data:"):
            picture.picture.path = _decode_data_uri(source, temp_dir) or ""
            continue
        local = resolve_app_path_to_filesystem(source)
        if local:
            picture.picture.path = local
        elif source.startswith(("http://", "https://")):
            to_fetch.setdefault(source, []).append(picture)
        else:
            picture.picture.path = ""
        picture.picture.is_network = False
    if not to_fetch:
        return
    own = _own_origins()
    async with httpx.AsyncClient(timeout=DOWNLOAD_TIMEOUT_S, follow_redirects=False) as client:
        urls = list(to_fetch)
        paths = await asyncio.gather(*(_download(client, url, temp_dir, own) for url in urls))
    for url, path in zip(urls, paths):
        for picture in to_fetch[url]:
            picture.picture.path = path or ""


# --------------------------------------------------------------------------- the writer


class PptxWriter:
    def __init__(self, model: PptxPresentationModel, temp_dir: str):
        self._model = model
        self._temp_dir = temp_dir
        self._ppt = Presentation()
        self._ppt.slide_width = Pt(px_to_pt(SLIDE_WIDTH_PX))
        self._ppt.slide_height = Pt(px_to_pt(SLIDE_HEIGHT_PX))

    def build(self) -> None:
        for slide_model in self._model.slides:
            self.add_slide(slide_model)

    def to_bytes(self) -> bytes:
        buffer = io.BytesIO()
        self._ppt.save(buffer)
        return finalize_package(buffer.getvalue())

    def save(self, path: str) -> None:
        data = self.to_bytes()
        with open(path, "wb") as f:
            f.write(data)

    # ----------------------------------------------------------------- slides and shapes

    def add_slide(self, slide_model: PptxSlideModel) -> None:
        slide = self._ppt.slides.add_slide(self._ppt.slide_layouts[BLANK_SLIDE_LAYOUT])
        if slide_model.background:
            self.apply_fill(slide.background, slide_model.background)
        if slide_model.note:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is not None:
                notes_frame.text = slide_model.note
        for shape_model in slide_model.shapes:
            if isinstance(shape_model, PptxPictureBoxModel):
                self.add_picture(slide, shape_model)
            elif isinstance(shape_model, PptxAutoShapeBoxModel):
                self.add_autoshape(slide, shape_model)
            elif isinstance(shape_model, PptxTextBoxModel):
                self.add_textbox(slide, shape_model)
            elif isinstance(shape_model, PptxConnectorModel):
                self.add_connector(slide, shape_model)

    def add_connector(self, slide: Slide, model: PptxConnectorModel) -> None:
        if model.thickness == 0:
            return
        connector = slide.shapes.add_connector(model.type, *model.position.to_pt_xyxy())
        if model.rotation:
            connector.rotation = model.rotation
        connector.line.width = Pt(model.thickness)
        connector.line.color.rgb = RGBColor.from_string(model.color)

    def add_picture(self, slide: Slide, model: PptxPictureBoxModel) -> None:
        image_path = model.picture.path
        if not image_path:
            return
        width, height = int(model.position.width), int(model.position.height)
        if model.clip or model.border_radius or model.invert or model.opacity or model.object_fit or model.shape:
            try:
                image: Image.Image = Image.open(image_path)
                image.load()
            except Exception:
                LOGGER.warning("pptx: could not open image %s", image_path)
                return
            untouched = image
            image = image.convert("RGBA")
            if model.border_radius:
                image = round_image_corners(image, [int(r) for r in model.border_radius])
            if model.object_fit:
                image = fit_image(image, width, height, model.object_fit)
            elif model.clip:
                image = clip_image(image, width, height)
            if model.border_radius:
                image = round_image_corners(image, [int(r) for r in model.border_radius])
            if model.shape == PptxBoxShapeEnum.CIRCLE:
                image = create_circle_image(image)
            if model.invert:
                image = invert_image(image)
            if model.opacity:
                image = set_image_opacity(image, model.opacity)
            if not self._unchanged_png(untouched, image):
                image_path = os.path.join(self._temp_dir, f"{uuid.uuid4()}.png")
                image.save(image_path, compress_level=1)
        position = self.margined_position(model.position, model.margin)
        try:
            picture = slide.shapes.add_picture(image_path, *position.to_pt_list())
        except Exception as exc:
            LOGGER.warning("pptx: could not add image %s (%s)", image_path, exc)
            return
        if model.rotation:
            picture.rotation = model.rotation

    @staticmethod
    def _unchanged_png(original: Image.Image, result: Image.Image) -> bool:
        """A PNG the steps above left as it was (a screenshot already at its box size):
        embed the original file instead of encoding it again."""
        if original.format != "PNG" or original.mode != "RGBA" or result.size != original.size:
            return False
        return result is original or result.tobytes() == original.tobytes()

    def add_autoshape(self, slide: Slide, model: PptxAutoShapeBoxModel) -> None:
        autoshape = slide.shapes.add_shape(model.type, *model.position.to_pt_list())
        if model.rotation:
            autoshape.rotation = model.rotation
        text_frame = autoshape.text_frame
        text_frame.word_wrap = model.text_wrap
        self.apply_fill(autoshape, model.fill)
        self.apply_margin(text_frame, model.margin)
        self.apply_stroke(autoshape, model.stroke)
        self.apply_shadow(autoshape, model.shadow)
        self.apply_border_radius(autoshape, model.border_radius)
        if model.paragraphs:
            self.add_paragraphs(text_frame, model.paragraphs)

    def add_textbox(self, slide: Slide, model: PptxTextBoxModel) -> None:
        textbox = slide.shapes.add_textbox(*model.position.to_pt_list())
        if model.rotation:
            textbox.rotation = model.rotation
        text_frame = textbox.text_frame
        text_frame.word_wrap = model.text_wrap
        if model.vertical_alignment:
            text_frame.vertical_anchor = model.vertical_alignment
        self.apply_fill(textbox, model.fill)
        self.apply_margin(text_frame, model.margin)
        self.apply_stroke(textbox, model.stroke)
        self.apply_shadow(textbox, model.shadow)
        self.add_paragraphs(text_frame, model.paragraphs)

    # ----------------------------------------------------------------- paragraphs

    def add_paragraphs(self, text_frame: TextFrame, paragraph_models: Iterable[PptxParagraphModel]) -> None:
        normalized = [p for model in paragraph_models for p in self.split_lines(model)]
        for index, paragraph_model in enumerate(normalized):
            paragraph = text_frame.add_paragraph() if index > 0 else text_frame.paragraphs[0]
            self.populate_paragraph(paragraph, paragraph_model)

    def split_lines(self, model: PptxParagraphModel) -> List[PptxParagraphModel]:
        """One paragraph per line: PowerPoint breaks lines between paragraphs."""
        runs = model.text_runs or []
        text = model.text or ""
        if not runs:
            if "\n" not in text:
                return [model]
            return [self.take_bullet_prefix(model, segment) for segment in text.split("\n")]
        segments: List[List[PptxTextRunModel]] = [[]]
        for run in runs:
            if "\n" not in run.text:
                segments[-1].append(run)
                continue
            parts = run.text.split("\n")
            for index, part in enumerate(parts):
                if part:
                    segments[-1].append(PptxTextRunModel(text=part, font=run.font))
                if index < len(parts) - 1:
                    segments.append([])
        return [
            self.take_bullet_prefix(model, "") if not segment else self.take_bullet_prefix(model, "", segment)
            for segment in segments
        ]

    @staticmethod
    def take_bullet_prefix(
        model: PptxParagraphModel, text: str, runs: Optional[List[PptxTextRunModel]] = None
    ) -> PptxParagraphModel:
        """A line typed as "• x", "- x" or "* x" becomes a bulleted paragraph."""
        bullet = model.bullet

        def strip(value: str) -> str:
            nonlocal bullet
            for prefix in BULLET_PREFIXES:
                if value.startswith(prefix) and not bullet:
                    bullet = PptxParagraphBulletModel(type="default")
                    return value[len(prefix) :]
            return value

        if runs is None:
            cleaned = strip(text)
            return model.model_copy(update={"text": cleaned, "text_runs": None, "bullet": bullet})
        if not runs:
            return model.model_copy(update={"text": "", "text_runs": None, "bullet": bullet})
        first = runs[0]
        cleaned = strip(first.text)
        if cleaned != first.text:
            runs = [PptxTextRunModel(text=cleaned, font=first.font), *runs[1:]]
        return model.model_copy(update={"text": None, "text_runs": runs, "bullet": bullet})

    def populate_paragraph(self, paragraph: _Paragraph, model: PptxParagraphModel) -> None:
        if model.spacing:
            self.apply_spacing(paragraph, model.spacing)
        if model.line_height and model.font and model.font.size:
            paragraph.line_spacing = (model.line_height / model.font.size) * 0.8
        if model.alignment:
            paragraph.alignment = model.alignment
        if model.letter_spacing:
            self.apply_letter_spacing(paragraph, model.letter_spacing)
        if model.font:
            self.apply_font(paragraph.font, model.font)
        if model.level:
            paragraph.level = model.level
        if model.bullet:
            self.apply_bullet(paragraph, model.bullet)
            paragraph.add_run().text = "  "
        text_runs: List[PptxTextRunModel] = []
        if model.text_runs:
            text_runs = model.text_runs
        elif model.text:
            text_runs = [PptxTextRunModel(text=model.text, font=model.font or PptxFontModel())]
        for run_model in text_runs:
            self.populate_run(paragraph.add_run(), run_model)
        if text_has_rtl(model.text) or any(text_has_rtl(run.text) for run in text_runs):
            paragraph._p.get_or_add_pPr().set("rtl", "1")

    def populate_run(self, run: _Run, model: PptxTextRunModel) -> None:
        run.text = model.text or ""
        if model.font:
            self.apply_font(run.font, model.font)

    # ----------------------------------------------------------------- formatting

    def apply_font(self, font: Font, model: PptxFontModel) -> None:
        font.name = model.name
        font.color.rgb = RGBColor.from_string(model.color)
        font.italic = model.italic
        font.size = Pt(px_to_pt(model.size))
        font.bold = (model.font_weight or 400) > 600
        if model.underline is not None:
            font.underline = bool(model.underline)
        if model.strike is not None:
            font._element.set("strike", "sngStrike" if model.strike else "noStrike")

    def apply_spacing(self, paragraph: _Paragraph, spacing: PptxSpacingModel) -> None:
        paragraph.space_before = Pt(gtz(px_to_pt(spacing.top)))
        paragraph.space_after = Pt(gtz(px_to_pt(spacing.bottom)))
        p_pr = paragraph._p.get_or_add_pPr()
        p_pr.set("marL", str(Pt(gtz(px_to_pt(spacing.left)))))
        p_pr.set("marR", str(Pt(gtz(px_to_pt(spacing.right)))))

    def apply_letter_spacing(self, paragraph: _Paragraph, letter_spacing: float) -> None:
        def_r_pr = paragraph._p.get_or_add_pPr().get_or_add_defRPr()
        def_r_pr.set("spc", str(int(px_to_pt(letter_spacing) * 100)))

    def apply_bullet(self, paragraph: _Paragraph, bullet: PptxParagraphBulletModel) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        for tag in ("buNone", "buChar", "buAutoNum", "buBlip", "buClr", "buSzPct", "buSzPts", "buFont"):
            for existing in p_pr.findall(f"{{{NS_A}}}{tag}"):
                p_pr.remove(existing)
        elements = []
        if bullet.color:
            color = etree.Element(f"{{{NS_A}}}buClr")
            etree.SubElement(color, f"{{{NS_A}}}srgbClr", {"val": bullet.color})
            elements.append(color)
        if bullet.type == "number":
            elements.append(etree.Element(f"{{{NS_A}}}buAutoNum", {"type": "arabicPeriod", "startAt": "1"}))
        else:
            elements.append(etree.Element(f"{{{NS_A}}}buChar", {"char": BULLET_CHAR}))
        # Schema order: colour, then the bullet, both before tabLst/defRPr/extLst.
        anchor = next((child for child in p_pr if etree.QName(child).localname in _AFTER_BULLETS), None)
        for element in elements:
            if anchor is None:
                p_pr.append(element)
            else:
                anchor.addprevious(element)

    def apply_fill(self, shape: object, fill: Optional[PptxFillModel]) -> None:
        shape_fill = shape.fill  # type: ignore[attr-defined]
        if not fill:
            shape_fill.background()
            return
        shape_fill.solid()
        shape_fill.fore_color.rgb = RGBColor.from_string(fill.color)
        self.set_fill_opacity(shape_fill, fill.opacity)

    def apply_stroke(self, shape: Shape, stroke: Optional[PptxStrokeModel]) -> None:
        if not stroke or stroke.thickness == 0:
            shape.line.fill.background()
            return
        shape.line.fill.solid()
        shape.line.fill.fore_color.rgb = RGBColor.from_string(stroke.color)
        shape.line.width = Pt(px_to_pt(stroke.thickness))
        self.set_fill_opacity(shape.line.fill, stroke.opacity)

    @staticmethod
    def set_fill_opacity(fill: object, opacity: Optional[float]) -> None:
        if opacity is None or opacity >= 1.0:
            return
        try:
            color = fill._xPr.solidFill.get_or_change_to_srgbClr()  # type: ignore[attr-defined]
        except AttributeError:
            return
        alpha = OxmlElement("a:alpha")
        alpha.set("val", str(int(opacity * 100000)))
        color.append(alpha)

    def apply_shadow(self, shape: Shape, shadow: Optional[PptxShadowModel]) -> None:
        sp_pr = shape._element.xpath("p:spPr")[0]
        effect_list = sp_pr.find(f"{{{NS_A}}}effectLst")
        if effect_list is None:
            effect_list = etree.SubElement(sp_pr, f"{{{NS_A}}}effectLst")
        else:
            for tag in ("outerShdw", "innerShdw", "prstShdw"):
                for old in effect_list.findall(f"{{{NS_A}}}{tag}"):
                    effect_list.remove(old)
        if shadow is None:
            # An explicit invisible shadow, so the theme never adds one.
            outer = etree.SubElement(effect_list, f"{{{NS_A}}}outerShdw", {"blurRad": "0", "dist": "0", "dir": "0"})
            color = etree.SubElement(outer, f"{{{NS_A}}}srgbClr", {"val": "000000"})
            etree.SubElement(color, f"{{{NS_A}}}alpha", {"val": "0"})
            return
        angle = int(round((shadow.angle % 360) * 60000)) if shadow.angle is not None else 0
        outer = etree.SubElement(
            effect_list,
            f"{{{NS_A}}}outerShdw",
            {
                "blurRad": f"{Pt(px_to_pt(shadow.radius))}",
                "dir": f"{angle}",
                "dist": f"{Pt(px_to_pt(shadow.offset))}",
                "rotWithShape": "0",
            },
        )
        color = etree.SubElement(outer, f"{{{NS_A}}}srgbClr", {"val": shadow.color})
        etree.SubElement(color, f"{{{NS_A}}}alpha", {"val": f"{int(shadow.opacity * 100000)}"})

    @staticmethod
    def apply_border_radius(shape: Shape, border_radius: Optional[float]) -> None:
        if not border_radius:
            return
        try:
            shape.adjustments[0] = Pt(px_to_pt(border_radius)) / min(shape.width, shape.height)
        except (IndexError, ZeroDivisionError):
            pass

    @staticmethod
    def apply_margin(text_frame: TextFrame, margin: Optional[PptxSpacingModel]) -> None:
        text_frame.margin_left = Pt(px_to_pt(margin.left) if margin else 0)
        text_frame.margin_right = Pt(px_to_pt(margin.right) if margin else 0)
        text_frame.margin_top = Pt(px_to_pt(margin.top) if margin else 0)
        text_frame.margin_bottom = Pt(px_to_pt(margin.bottom) if margin else 0)

    @staticmethod
    def margined_position(position: PptxPositionModel, margin: Optional[PptxSpacingModel]) -> PptxPositionModel:
        if not margin:
            return position
        return PptxPositionModel(
            left=position.left + margin.left,
            top=position.top + margin.top,
            width=max(position.width - margin.left - margin.right, 0),
            height=max(position.height - margin.top - margin.bottom, 0),
        )


# --------------------------------------------------------------------------- package fixes


def _xml_bytes(root: etree._Element) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _slide_with_group_xfrm(xml: bytes) -> bytes:
    """Keynote rejects a group shape without a:xfrm (python-pptx writes an empty grpSpPr)."""
    root = etree.fromstring(xml)
    changed = False
    for grp_sp_pr in root.iter(f"{{{NS_P}}}grpSpPr"):
        if grp_sp_pr.find(f"{{{NS_A}}}xfrm") is None:
            xfrm = etree.SubElement(grp_sp_pr, f"{{{NS_A}}}xfrm")
            etree.SubElement(xfrm, f"{{{NS_A}}}off", x="0", y="0")
            etree.SubElement(xfrm, f"{{{NS_A}}}ext", cx="0", cy="0")
            etree.SubElement(xfrm, f"{{{NS_A}}}chOff", x="0", y="0")
            etree.SubElement(xfrm, f"{{{NS_A}}}chExt", cx="0", cy="0")
            changed = True
    return _xml_bytes(root) if changed else xml


def _notes_master_fixes(parts: dict[str, bytes]) -> None:
    """List the notes master in presentation.xml (python-pptx only relates it)."""
    rels_name, pres_name = "ppt/_rels/presentation.xml.rels", "ppt/presentation.xml"
    if rels_name not in parts or pres_name not in parts:
        return
    if not any(re.fullmatch(r"ppt/notesMasters/[^/]+\.xml", name) for name in parts):
        return
    rels = etree.fromstring(parts[rels_name])
    rel_tag = f"{{{NS_PKG_REL}}}Relationship"
    ids = {rel.get("Id") for rel in rels.findall(rel_tag)}
    notes_rel = next((rel for rel in rels.findall(rel_tag) if rel.get("Type") == NOTES_MASTER_REL), None)
    if notes_rel is None:
        next_id = 1
        while f"rId{next_id}" in ids:
            next_id += 1
        notes_rel = etree.SubElement(rels, rel_tag, Id=f"rId{next_id}", Type=NOTES_MASTER_REL)
        notes_rel.set("Target", "notesMasters/notesMaster1.xml")
        parts[rels_name] = _xml_bytes(rels)
    pres = etree.fromstring(parts[pres_name])
    id_list = pres.find(f"{{{NS_P}}}notesMasterIdLst")
    if id_list is None:
        id_list = etree.Element(f"{{{NS_P}}}notesMasterIdLst")
        masters = pres.find(f"{{{NS_P}}}sldMasterIdLst")
        pres.insert(list(pres).index(masters) + 1 if masters is not None else 0, id_list)
    if not id_list.findall(f"{{{NS_P}}}notesMasterId"):
        etree.SubElement(id_list, f"{{{NS_P}}}notesMasterId", {f"{{{NS_R}}}id": notes_rel.get("Id") or ""})
        parts[pres_name] = _xml_bytes(pres)


def _without_thumbnail(rels_xml: bytes) -> bytes:
    rels = etree.fromstring(rels_xml)
    removed = False
    for rel in rels.findall(f"{{{NS_PKG_REL}}}Relationship"):
        if rel.get("Type") == THUMBNAIL_REL or rel.get("Target") == THUMBNAIL_PART:
            rels.remove(rel)
            removed = True
    return _xml_bytes(rels) if removed else rels_xml


def finalize_package(data: bytes) -> bytes:
    """Compatibility fixes on the saved package, in memory: group xfrm for Keynote, the
    notes master listed, python-pptx's template thumbnail dropped."""
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        infos = source.infolist()
        parts = {info.filename: source.read(info.filename) for info in infos}
    for name in list(parts):
        if re.fullmatch(r"ppt/slides/slide\d+\.xml", name):
            parts[name] = _slide_with_group_xfrm(parts[name])
    _notes_master_fixes(parts)
    if "_rels/.rels" in parts:
        parts["_rels/.rels"] = _without_thumbnail(parts["_rels/.rels"])
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in infos:
            if info.filename == THUMBNAIL_PART:
                continue
            target.writestr(info.filename, parts[info.filename])
    return out.getvalue()


# --------------------------------------------------------------------------- entry points


def write_pptx(model: PptxPresentationModel, path: str, temp_dir: str) -> None:
    """Build and save the deck (pictures must already point at local files)."""
    writer = PptxWriter(model, temp_dir)
    writer.build()
    writer.save(path)


async def write_pptx_to_exports(model: PptxPresentationModel) -> str:
    """Write the deck into <data>/exports and return the file's path."""
    from utils.asset_directory_utils import get_exports_directory
    from utils.get_env import get_temp_directory_env

    temp_root = get_temp_directory_env() or os.path.join(tempfile.gettempdir(), "chatppt")
    os.makedirs(temp_root, exist_ok=True)
    out_path = os.path.join(get_exports_directory(), f"{export_basename(model.name)}_{uuid.uuid4()}.pptx")
    with tempfile.TemporaryDirectory(prefix="pptx-", dir=temp_root) as temp_dir:
        await resolve_picture_sources(model, temp_dir)
        await asyncio.to_thread(write_pptx, model, out_path, temp_dir)
    os.chmod(out_path, 0o644)
    return out_path

